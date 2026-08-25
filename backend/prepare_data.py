"""
Week 2 data collection helper.

Drop raw course materials into data/raw/ using this naming convention:
    <topic-slug>__<original-filename>.pdf/.pptx/.docx

Example:
    er-diagrams__lecture3_slides.pptx
    normalization__week4_notes.pdf

Run:
    python prepare_data.py

Output:
    data/processed/<topic-slug>/<filename>.txt   (cleaned text)
    data/processed/manifest.csv                  (index of everything processed)
"""

import os
import re
import csv
import pathlib

# pyrefly: ignore [missing-import]
import pdfplumber
# pyrefly: ignore [missing-import]
from pptx import Presentation
# pyrefly: ignore [missing-import]
from docx import Document

RAW_DIR = pathlib.Path(__file__).parent.parent / "data" / "raw"
PROCESSED_DIR = pathlib.Path(__file__).parent.parent / "data" / "processed"


def clean_text(text: str) -> str:
    """Basic cleanup: collapse whitespace, drop obvious page-number-only lines."""
    lines = text.split("\n")
    cleaned = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if re.fullmatch(r"\d{1,3}", stripped):  # bare page number
            continue
        if re.fullmatch(r"page\s*\d+", stripped, re.IGNORECASE):
            continue
        cleaned.append(stripped)
    text = "\n".join(cleaned)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: pathlib.Path) -> str:
    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_pptx(path: pathlib.Path) -> str:
    try:
        prs = Presentation(path)
        text_parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line.strip():
                            text_parts.append(line)
        return "\n".join(text_parts)
    except Exception:
        import zipfile
        text_parts = []
        with zipfile.ZipFile(path, 'r') as z:
            slide_files = sorted([f for f in z.namelist() if f.startswith('ppt/slides/slide') and f.endswith('.xml')],
                                 key=lambda x: int(re.search(r'\d+', x).group()) if re.search(r'\d+', x) else 0)
            for slide_file in slide_files:
                xml_content = z.read(slide_file).decode('utf-8', errors='ignore')
                matches = re.findall(r'<a:t[^>]*>(.*?)</a:t>', xml_content, re.DOTALL)
                slide_text = ' '.join([m.strip() for m in matches if m.strip()])
                if slide_text:
                    text_parts.append(slide_text)
        return "\n\n".join(text_parts)


def extract_docx(path: pathlib.Path) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_html(path: pathlib.Path) -> str:
    content = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r'<script.*?>.*?</script>', '', content, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r'<style.*?>.*?</style>', '', text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r'<[^>]+>', ' ', text)
    return text


def extract_ppt(path: pathlib.Path) -> str:
    content = path.read_bytes()
    ascii_strings = re.findall(rb'[\x20-\x7E]{4,}', content)
    utf16_strings = re.findall(rb'(?:[\x20-\x7E]\x00){4,}', content)
    decoded_ascii = [s.decode('ascii', errors='ignore') for s in ascii_strings]
    decoded_utf16 = [s.decode('utf-8', errors='ignore') for s in utf16_strings]
    return "\n".join(decoded_ascii + decoded_utf16)


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".docx": extract_docx,
    ".html": extract_html,
    ".htm": extract_html,
    ".ppt": extract_ppt,
}


def main():
    all_files = sorted([p for p in RAW_DIR.rglob("*") if p.is_file() and p.name != ".gitkeep"])
    if not all_files:
        print(f"No files found in {RAW_DIR}. Add course materials there first.")
        return

    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    manifest_rows = []

    for file_path in all_files:
        if file_path.suffix.lower() not in EXTRACTORS:
            print(f"Skipping unsupported file type: {file_path.name}")
            continue

        if "__" in file_path.stem:
            topic_slug, original_name = file_path.stem.split("__", 1)
        elif file_path.parent != RAW_DIR:
            topic_slug = re.sub(r'[^a-zA-Z0-9]+', '-', file_path.parent.name).strip('-').lower()
            original_name = re.sub(r'[^a-zA-Z0-9_-]+', '_', file_path.stem)
        else:
            topic_slug = "uncategorized"
            original_name = file_path.stem

        extractor = EXTRACTORS[file_path.suffix.lower()]

        try:
            raw_text = extractor(file_path)
        except Exception as e:
            print(f"FAILED to extract {file_path.name}: {e}")
            continue

        cleaned = clean_text(raw_text)

        topic_dir = PROCESSED_DIR / topic_slug
        topic_dir.mkdir(parents=True, exist_ok=True)
        out_path = topic_dir / f"{original_name}.txt"
        out_path.write_text(cleaned, encoding="utf-8")

        word_count = len(cleaned.split())
        manifest_rows.append({
            "source_file": file_path.name,
            "topic": topic_slug,
            "output_path": str(out_path.relative_to(PROCESSED_DIR.parent.parent)),
            "word_count": word_count,
        })
        print(f"Processed: {file_path.name} -> {out_path.name} ({word_count} words, topic={topic_slug})")

    manifest_path = PROCESSED_DIR / "manifest.csv"
    with open(manifest_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["source_file", "topic", "output_path", "word_count"])
        writer.writeheader()
        writer.writerows(manifest_rows)

    print(f"\nDone. {len(manifest_rows)} files processed.")
    print(f"Manifest written to {manifest_path}")


if __name__ == "__main__":
    main()
